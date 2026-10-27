"""Deterministic in-memory document builders for tests (no binary fixtures in the repo)."""

from __future__ import annotations

import io
import zipfile

from docx import Document
from pypdf import PdfReader, PdfWriter


def _pdf_escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf(pages: list[list[str]], title: str | None = None) -> bytes:
    """Build a valid text PDF. Each page is a list of lines; ``""`` produces a blank line."""
    objects: list[bytes] = []

    def add(body: bytes) -> int:
        objects.append(body)
        return len(objects)

    font_id = add(
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>"
    )
    pages_id_placeholder = len(objects) + 1
    objects.append(b"")  # pages tree, filled in below
    page_ids: list[int] = []
    for lines in pages:
        ops = ["BT", "/F1 11 Tf", "14 TL", "72 760 Td"]
        for line in lines:
            if line:
                ops.append(f"({_pdf_escape(line)}) Tj")
            ops.append("T*")
        ops.append("ET")
        stream = "\n".join(ops).encode("latin-1")
        content_id = add(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
        page_ids.append(
            add(
                b"<< /Type /Page /Parent %d 0 R /MediaBox [0 0 612 792] "
                b"/Resources << /Font << /F1 %d 0 R >> >> /Contents %d 0 R >>"
                % (pages_id_placeholder, font_id, content_id)
            )
        )
    kids = b" ".join(b"%d 0 R" % pid for pid in page_ids)
    objects[pages_id_placeholder - 1] = b"<< /Type /Pages /Kids [%s] /Count %d >>" % (
        kids,
        len(page_ids),
    )
    catalog_id = add(b"<< /Type /Catalog /Pages %d 0 R >>" % pages_id_placeholder)
    info_id = add(b"<< /Title (%s) >>" % _pdf_escape(title).encode("latin-1")) if title else None

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref_at = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    trailer = b"<< /Size %d /Root %d 0 R" % (len(objects) + 1, catalog_id)
    if info_id:
        trailer += b" /Info %d 0 R" % info_id
    out.write(b"trailer\n" + trailer + b" >>\nstartxref\n%d\n%%%%EOF\n" % xref_at)
    return out.getvalue()


def encrypt_pdf(pdf: bytes, user_password: str, owner_password: str = "owner") -> bytes:
    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(pdf)))
    writer.encrypt(user_password=user_password, owner_password=owner_password, algorithm="RC4-128")
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def make_docx(build: list[tuple[str, object]], title: str | None = None) -> bytes:
    """Build a DOCX from ``(kind, payload)`` steps.

    Kinds: ``("heading", (text, level))``, ``("paragraph", text)``, ``("bullet", text)``,
    ``("table", [[cell, ...], ...])``.
    """
    document = Document()
    if title:
        document.core_properties.title = title
    for kind, payload in build:
        if kind == "heading":
            text, level = payload  # type: ignore[misc]
            document.add_heading(text, level=level)
        elif kind == "paragraph":
            document.add_paragraph(str(payload))
        elif kind == "bullet":
            document.add_paragraph(str(payload), style="List Bullet")
        elif kind == "table":
            rows: list[list[str]] = payload  # type: ignore[assignment]
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    table.cell(r, c).text = value
        else:
            raise ValueError(kind)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def make_zip(entries: dict[str, bytes]) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, content in entries.items():
            archive.writestr(name, content)
    return out.getvalue()
